"""Build the REV-ECIT manuscript from the official Word template."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "report" / "REV_ECIT_2026_manuscript.md"
TEMPLATE = ROOT / "docs" / "report" / "REV-ECIT_official_template.docx"
OUTPUT = ROOT / "docs" / "report" / "SecureEdu_REV_ECIT_2026.docx"


def _set_columns(section, count: int, space_twips: int = 360) -> None:
    section_properties = section._sectPr
    columns = section_properties.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        section_properties.append(columns)
    columns.set(qn("w:num"), str(count))
    columns.set(qn("w:space"), str(space_twips))


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 35, start: int = 45, bottom: int = 35, end: int = 45) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "808080")


def _plain(text: str) -> str:
    return text.replace("**", "").replace("`", "").replace("*", "")


def _add_rich_text(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(8)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def _configure_document(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Inches(0.17)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    first = document.sections[0]
    first.page_width = Cm(21.0)
    first.page_height = Cm(29.7)
    first.top_margin = Inches(0.72)
    first.bottom_margin = Inches(0.72)
    first.left_margin = Inches(0.78)
    first.right_margin = Inches(0.78)
    first.header_distance = Inches(0.3)
    first.footer_distance = Inches(0.3)
    _set_columns(first, 1)

    document.core_properties.title = (
        "SecureEdu: Keyed Audit Chaining and External Checkpoints for "
        "Versioned Student Records"
    )
    document.core_properties.author = "Trần Thị Hà Vy"
    document.core_properties.subject = "REV-ECIT 2026 manuscript"


def _add_front_matter(document: Document, lines: list[str]) -> int:
    title = lines[0].removeprefix("# ").strip()
    paragraph = document.add_paragraph(style="paper title")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(title)

    index = 1
    front: list[str] = []
    while index < len(lines) and not lines[index].startswith("## Abstract"):
        if lines[index].strip():
            front.append(_plain(lines[index].strip().removesuffix("  ")))
        index += 1
    for position, text in enumerate(front):
        style = "Author" if position == 0 else "Affiliation"
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.space_before = Pt(2 if position == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        if text.startswith("["):
            run.bold = True
            run.font.color.rgb = None

    index += 1
    while index < len(lines) and not lines[index].strip():
        index += 1
    abstract = lines[index].strip()
    paragraph = document.add_paragraph(style="Abstract")
    paragraph.paragraph_format.space_after = Pt(3)
    lead = paragraph.add_run("Abstract—")
    lead.italic = True
    paragraph.add_run(abstract)
    index += 1

    while index < len(lines) and not lines[index].strip():
        index += 1
    keywords = _plain(lines[index].strip())
    paragraph = document.add_paragraph(style="key words")
    paragraph.paragraph_format.space_after = Pt(4)
    _add_rich_text(paragraph, keywords)
    return index + 1


def _configure_body_section(section, columns: int) -> None:
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    _set_columns(section, columns)


def _add_table(document: Document, rows: list[list[str]], caption_text: str) -> None:
    wide_section = document.add_section(WD_SECTION.CONTINUOUS)
    _configure_body_section(wide_section, 1)
    caption = document.add_paragraph()
    _add_rich_text(caption, caption_text)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_after = Pt(2)
    for run in caption.runs:
        run.bold = True
        run.font.size = Pt(7.5)

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Normal Table"
    _set_table_borders(table)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            cell.text = _plain(value)
            if row_index == 0:
                _set_cell_shading(cell, "D9E2F3")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = None
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 0.9
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if column_index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.keep_with_next = row_index < len(rows) - 1
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(7.5)
                    run.bold = row_index == 0
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.first_line_indent = None
    body_section = document.add_section(WD_SECTION.CONTINUOUS)
    _configure_body_section(body_section, 2)


def _add_body(document: Document, lines: list[str], start: int) -> None:
    body_section = document.add_section(WD_SECTION.CONTINUOUS)
    _configure_body_section(body_section, 2)

    index = start
    pending_caption = ""
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            heading_text = _plain(line[3:])
            unnumbered = heading_text in {"ACKNOWLEDGMENT", "REFERENCES"}
            style = "Heading 5" if unnumbered else "Heading 1"
            heading = document.add_paragraph(style=style)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            heading.paragraph_format.space_before = Pt(5)
            heading.paragraph_format.space_after = Pt(2)
            heading.paragraph_format.keep_with_next = True
            if not unnumbered:
                heading_text = re.sub(r"^[IVX]+\.\s*", "", heading_text)
            run = heading.add_run(heading_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.bold = False
            index += 1
            continue
        if line.startswith("### "):
            heading = document.add_paragraph()
            heading.paragraph_format.first_line_indent = None
            heading.paragraph_format.space_before = Pt(3)
            heading.paragraph_format.space_after = Pt(1)
            heading.paragraph_format.keep_with_next = True
            run = heading.add_run(_plain(line[4:]))
            run.bold = True
            run.italic = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            index += 1
            continue
        if line.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            _add_table(document, rows, pending_caption)
            pending_caption = ""
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.widow_control = True
        if line.startswith("[") and re.match(r"\[\d+\]", line):
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.paragraph_format.line_spacing = 0.92
            _add_rich_text(paragraph, line)
            for run in paragraph.runs:
                run.font.size = Pt(7.5)
        elif line.startswith("**Table "):
            document._element.body.remove(paragraph._element)
            pending_caption = line
        else:
            _add_rich_text(paragraph, line)
        index += 1


def main() -> int:
    retained_hash = __import__("hashlib").sha256(TEMPLATE.read_bytes()).hexdigest().upper()
    expected = "19DCD2E8A6183A2674961ADACB58C784B52A1FF0BF364C60D0CD16929E18EC60"
    if retained_hash != expected:
        raise RuntimeError("The official template changed; distill it again before building.")

    shutil.copy2(TEMPLATE, OUTPUT)
    document = Document(OUTPUT)
    _configure_document(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = _add_front_matter(document, lines)
    _add_body(document, lines, start)
    document.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
